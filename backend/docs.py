"""Reads PDF and Word files, and turns Markdown answers into Word documents."""
import base64
import binascii
import io
import re


def _pdf_pages_as_images(raw: bytes, max_pages: int = 5) -> list[str]:
    """For scanned PDFs: render the first pages as pictures so a vision model can read them."""
    import pymupdf
    out = []
    with pymupdf.open(stream=raw, filetype="pdf") as doc:
        for page in list(doc)[:max_pages]:
            pix = page.get_pixmap(matrix=pymupdf.Matrix(1.4, 1.4))
            out.append("data:image/jpeg;base64," + base64.b64encode(pix.tobytes("jpeg")).decode())
    return out


def extract(name: str, data_b64: str) -> dict:
    """-> {"text": str, "images": [data URLs]} (images only for scanned PDFs)."""
    try:
        raw = base64.b64decode(data_b64, validate=True)
    except (binascii.Error, ValueError):
        raise ValueError("The file could not be read.")
    if len(raw) > 6_000_000:
        raise ValueError("That document is larger than 6 MB.")
    lower = name.lower()
    images: list[str] = []
    if lower.endswith(".pdf"):
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(raw))
        if reader.is_encrypted:
            raise ValueError("This PDF is password protected.")
        text = "\n".join((p.extract_text() or "") for p in reader.pages[:60]).strip()
        if len(text) < 150:  # scanned or image-only PDF
            try:
                images = _pdf_pages_as_images(raw)
            except Exception:
                images = []
    elif lower.endswith(".docx"):
        import docx
        d = docx.Document(io.BytesIO(raw))
        parts = [p.text for p in d.paragraphs]
        for table in d.tables:
            for row in table.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        text = "\n".join(parts).strip()
    else:
        raise ValueError("Only PDF and Word (.docx) documents are supported.")
    if not text and not images:
        raise ValueError("No readable content found in that document.")
    return {"text": text[:30000], "images": images}


_INLINE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\*[^*\n]+\*)")


def _runs(par, text):
    from docx.shared import Pt
    for chunk in _INLINE.split(text):
        if not chunk:
            continue
        if chunk.startswith("**") and chunk.endswith("**") and len(chunk) > 4:
            par.add_run(chunk[2:-2]).bold = True
        elif chunk.startswith("`") and chunk.endswith("`") and len(chunk) > 2:
            r = par.add_run(chunk[1:-1])
            r.font.name = "Consolas"
            r.font.size = Pt(10)
        elif chunk.startswith("*") and chunk.endswith("*") and len(chunk) > 2:
            par.add_run(chunk[1:-1]).italic = True
        else:
            par.add_run(chunk)


def markdown_to_docx(title: str, md: str) -> bytes:
    import docx
    from docx.shared import Pt
    d = docx.Document()
    d.core_properties.title = title[:120]
    lines = md.replace("\r\n", "\n").split("\n")
    if not any(l.startswith("# ") for l in lines[:3]):
        d.add_heading(title, level=0)
    i = 0
    while i < len(lines):
        line = lines[i]
        s = line.strip()
        if s.startswith("```"):
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                p = d.add_paragraph()
                r = p.add_run(lines[i])
                r.font.name = "Consolas"
                r.font.size = Pt(10)
                i += 1
            i += 1
            continue
        if s.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c):
                    rows.append(cells)
                i += 1
            if rows:
                cols = max(len(r) for r in rows)
                t = d.add_table(rows=len(rows), cols=cols)
                t.style = "Table Grid"
                for ri, r in enumerate(rows):
                    for ci in range(cols):
                        cell = t.cell(ri, ci)
                        cell.text = ""
                        _runs(cell.paragraphs[0], r[ci] if ci < len(r) else "")
                        if ri == 0:
                            for run in cell.paragraphs[0].runs:
                                run.bold = True
            continue
        m = re.match(r"^(#{1,4})\s+(.*)", s)
        if m:
            d.add_heading(m.group(2), level=len(m.group(1)))
        elif re.match(r"^[-*]\s+", s):
            _runs(d.add_paragraph(style="List Bullet"), re.sub(r"^[-*]\s+", "", s))
        elif re.match(r"^\d+[.)]\s+", s):
            _runs(d.add_paragraph(style="List Number"), re.sub(r"^\d+[.)]\s+", "", s))
        elif re.fullmatch(r"[-*_]{3,}", s):
            pass
        elif s:
            _runs(d.add_paragraph(), s)
        i += 1
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()